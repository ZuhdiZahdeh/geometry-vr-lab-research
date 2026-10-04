#!/usr/bin/env python3
"""Reproducibly transcribe the supplied research DOCX into site source content.

The July 2026 source is preserved as a dated document. Later developments belong
to their own pages; this transcription does not silently revise research claims.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.oxml.ns import qn


NOTICE = (
    "هذه الصفحة تنقل نص الإصدار الأكاديمي المرفق، المؤرّخ في 29 تموز/يوليو 2026. "
    "تصف حالات الإنجاز والإتاحة كما وردت في تلك النسخة، ولا تُعد وصفًا تلقائيًا "
    "للإصدار الحالي. تُوثّق التطويرات اللاحقة في صفحة مستقلة."
)


def safe_url(value: str) -> str | None:
    """Only publish ordinary web links; DOCX relationships can be file links."""
    value = value.strip()
    return value if urlparse(value).scheme.lower() in {"http", "https"} else None


def website_text(value: str) -> str:
    """Apply the website's requested Arabic punctuation style to prose."""
    return value.replace("\u061b", "\u060c")


def text_node(node) -> str:
    result = []
    for child in node.iter():
        if child.tag == qn("w:t"):
            result.append(child.text or "")
        elif child.tag == qn("w:tab"):
            result.append("\t")
        elif child.tag in {qn("w:br"), qn("w:cr")}:
            result.append("\n")
    return website_text("".join(result))


def link_plain_urls(text: str) -> str:
    result = []
    last = 0
    for match in re.finditer(r"https?://[^\s<>]+", text):
        raw = match.group().rstrip(".,;،؛)")
        result.append(html.escape(text[last:match.start()]))
        result.append('<a href="' + html.escape(raw, quote=True) + '" rel="noopener noreferrer">' + html.escape(raw) + "</a>")
        last = match.start() + len(raw)
    result.append(html.escape(text[last:]))
    return "".join(result).replace("\n", "<br>").replace("\t", " ")


def run_html(run, autolink=True) -> str:
    value = text_node(run)
    result = link_plain_urls(value) if autolink else html.escape(value).replace("\n", "<br>")
    if not result:
        return ""
    props = run.find(qn("w:rPr"))
    if props is not None:
        def enabled(name):
            element = props.find(qn("w:" + name))
            return element is not None and element.get(qn("w:val"), "true") not in {"false", "0", "off"}
        if enabled("i"):
            result = "<em>" + result + "</em>"
        if enabled("b"):
            result = "<strong>" + result + "</strong>"
        vert = props.find(qn("w:vertAlign"))
        if vert is not None and vert.get(qn("w:val")) in {"superscript", "subscript"}:
            tag = "sup" if vert.get(qn("w:val")) == "superscript" else "sub"
            result = "<" + tag + ">" + result + "</" + tag + ">"
    return result


def paragraph_html(element, rels) -> str:
    result = []
    for child in element:
        if child.tag == qn("w:r"):
            result.append(run_html(child))
        elif child.tag == qn("w:hyperlink"):
            rid = child.get(qn("r:id"))
            content = "".join(run_html(run, False) for run in child if run.tag == qn("w:r"))
            link = safe_url(str(rels[rid].target_ref)) if rid in rels else None
            if link:
                result.append('<a href="' + html.escape(link, quote=True) + '" rel="noopener noreferrer">' + content + "</a>")
            else:
                result.append(content)
        elif child.tag == qn("w:sdt"):
            # No SDTs occur in this source; preserve text if later versions have one.
            result.append(link_plain_urls(text_node(child)))
    return "".join(result)


def page_metadata(path: str, headings: list[dict]) -> tuple[str, str]:
    top = next((item for item in headings if item["path"] == path and item["anchor"] == "top"), None)
    if top:
        return top["title"], "chapter" if path.startswith("chapters/") else "appendix" if "appendix-" in path else "references" if path.startswith("references/") else "appendices"
    return {
        "about/index.html": ("عن البحث والباحث", "about"),
        "abstract/index.html": ("الملخص العربي والإنجليزي", "abstract"),
        "index.html": ("فهرس البحث في النسخة المصدرية", "source-contents"),
        "evidence/index.html": ("قائمة الجداول في البحث", "source-evidence"),
    }[path]


def write_markdown(page: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = ["---", "title: " + json.dumps(page["title"], ensure_ascii=False), "path: " + page["path"], "source_date: 2026-07-29", "kind: " + page["kind"], "---", "", "> " + page["notice"], ""]
    for block in page["blocks"]:
        lines.append("<!-- source: " + block["source"] + " -->")
        if block["type"] == "heading":
            lines.extend(["#" * max(1, min(6, block["level"])) + " " + block["text"] + " {#" + block["id"] + "}", ""])
        elif block["type"] == "paragraph":
            lines.extend(['<a id="' + block["id"] + '"></a>', block["html"], ""])
        elif block["type"] == "table":
            lines.extend(['<a id="' + block["id"] + '"></a>', "<!-- table role: " + block["role"] + " -->"])
            # HTML preserves multi-line cells without changing their content.
            lines.extend(["<table>", "<tbody>"])
            for row_index, row in enumerate(block["rows"]):
                tag = "th" if row_index == 0 and block["presentation"] == "table" else "td"
                lines.append("<tr>" + "".join("<" + tag + ">" + link_plain_urls(cell) + "</" + tag + ">" for cell in row) + "</tr>")
            lines.extend(["</tbody>", "</table>", ""])
        elif block["type"] == "image":
            lines.extend(['<a id="' + block["id"] + '"></a>', "![" + block["alt"] + "](../../" + block["src"] + ")", "", "*" + block["caption"] + "*", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def extract(source: Path, inventory: Path, destination: Path) -> dict:
    original = json.loads(inventory.read_text(encoding="utf-8"))
    doc = Document(source)
    mapping = {item["source_start"]: item for item in original["map"]}
    references = {item["source"]: item for item in original["references"]}
    tables = {item["id"]: item for item in original["tables"]}
    placements = {item["member"].split("/")[-1]: item for item in original["images"]}
    source_blocks = {item["locator"]: item for item in original["blocks"]}
    images_dir = destination / "assets/images"
    images_dir.mkdir(parents=True, exist_ok=True)
    figure_text = {
        "image1.png": ("ملف APK الناتج عن البناء في Unity في مستكشف ملفات Windows", "شكل (ط-1): ملف APK الناتج عن عملية البناء في Unity، بحجم يقارب 124 MB.", True),
        "image2.png": ("واجهة Meta Quest Developer Hub تعرض تطبيق Geometry VR Lab بالإصدار 1.0.0 على Quest 3S", "شكل (ط-2): ظهور التطبيق بالإصدار 1.0.0 وحالة Running على نظارة Quest 3S داخل Meta Quest Developer Hub.", True),
        "image3.jpeg": ("الأسطوانة وشاشة اختيار المجسمات في مختبر المجسمات السابق", "شكل (ط-3): اختيار الأسطوانة وعرض معلوماتها.", False),
        "image4.jpeg": ("الكرة ودليل سطحها وشاشة اختيار المجسمات في مختبر المجسمات السابق", "شكل (ط-4): اختيار الكرة وإظهار دليل هندسي حولها.", False),
        "image5.jpeg": ("مكعب تظهر عليه أدلة الوجه والحافة والرأس بألوان مختلفة", "شكل (ط-5): استكشاف المكعب وإظهار الوجه والحافة والرأس بألوان مختلفة.", False),
        "image6.jpeg": ("مكعب مع تسميات الوجه والحافة والرأس بعد تغيير زاوية العرض", "شكل (ط-6): المكعب بعد تغيير زاوية الرؤية، مع استمرار ظهور التسميات المكانية.", False),
        "image7.png": ("منصة المجسمات ولوحة مهام التهيئة وزر إعادة الأجسام في غرفة التدريب السابقة", "شكل (ط-7): لوحة مهام التدريب داخل غرفة التهيئة، وتظهر المجسمات وزر RESET OBJECTS قبل الانتقال إلى مختبر المجسمات.", False),
        "image8.png": ("إمساك الأسطوانة بوحدة التحكم داخل غرفة التدريب السابقة", "شكل (ط-8): تفاعل المستخدم مع الأسطوانة باستخدام وحدة التحكم، بما يوثّق عمل الإمساك والتحريك داخل التطبيق المستقل.", False),
        "image9.png": ("إمساك مكعب وشعاع التفاعل من وحدة التحكم داخل غرفة التدريب السابقة", "شكل (ط-9): الإمساك بالمكعب وتغيير موضعه داخل غرفة التهيئة، مع ظهور شعاع التفاعل من وحدة التحكم.", False),
    }
    figure_by_part = {}
    for name, item in placements.items():
        number = int(re.search(r"image(\d+)", name)[1])
        asset = "assets/images/figure-i-" + f"{number:02d}" + Path(name).suffix
        part = next((rel.target_part for rel in doc.part.rels.values() if not rel.is_external and str(rel.target_part.partname).endswith("/" + name)), None)
        if part is None:
            raise ValueError("Missing embedded image " + name)
        (destination / asset).write_bytes(part.blob)
        alt, caption, review = figure_text[name]
        figure_by_part[name] = {
            "type": "image", "id": "figure-i-" + f"{number:02d}", "src": asset,
            "alt": alt, "caption": caption, "source_image": item["image_id"],
            "sha256": hashlib.sha256(part.blob).hexdigest(), "width": item["width"], "height": item["height"],
            "publication_review": "فحص معلومات الجهاز والحساب والمسارات المحلية قبل النشر" if review else "لقطة تاريخية من النسخة السابقة",
            "requires_privacy_review": review,
        }
        if number in {3, 4}:
            figure_by_part[name]["caption_source"] = {"table": "T022", "row": 1, "column": number - 3}
        else:
            figure_by_part[name]["caption_source"] = {"paragraph": {1: "P0422", 2: "P0424", 5: "P0465", 6: "P0470", 7: "P0479", 8: "P0481", 9: "P0484"}[number]}
    pages = {}
    nonempty_paragraphs = []
    emitted_tables = []
    emitted_images = []
    empty_paragraphs = []
    pi = ti = 0
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            locator = f"P{pi + 1:04d}"
            paragraph = doc.paragraphs[pi]
            pi += 1
            source_block = source_blocks[locator]
            text = website_text(paragraph.text)
            path = source_block["target"].split("#")[0]
            if path not in pages:
                title, kind = page_metadata(path, original["map"])
                pages[path] = {"path": path, "title": title, "kind": kind, "notice": NOTICE, "source_date": "2026-07-29", "blocks": []}
            page = pages[path]
            if text.strip():
                nonempty_paragraphs.append(locator)
                heading = mapping.get(locator)
                ref = references.get(locator)
                block = {
                    "type": "heading" if heading else "paragraph", "source": locator,
                    "text": text, "html": paragraph_html(element, doc.part.rels),
                    "id": heading["anchor"] if heading else ref["id"].lower() if ref else "source-" + locator.lower(),
                    "style": paragraph.style.name,
                }
                if heading:
                    # Heading labels are exact source text. The map's title is a
                    # navigation label, e.g. the split normal-style appendix title.
                    level = 1 if heading["anchor"] == "top" else max(2, heading["level"])
                    block.update({"level": level, "navigation_title": heading["title"], "map_id": heading["id"], "parent_id": heading["parent_anchor"]})
                if ref:
                    block.update({"reference_id": ref["id"], "reference_kind": ref["kind"], "reference_url": ref["url"], "reference_status": ref["status"], "direction": "ltr"})
                elif path == "abstract/index.html" and locator >= "P0022":
                    block["direction"] = "ltr"
                if "List Bullet" in paragraph.style.name or text.lstrip().startswith("•"):
                    block["list_kind"] = "bullet"
                elif "List Number" in paragraph.style.name:
                    block["list_kind"] = "ordered"
                if paragraph.style.name.startswith("Caption") or text.startswith("شكل ("):
                    block["is_caption"] = True
                page["blocks"].append(block)
            else:
                empty_paragraphs.append(locator)
        elif element.tag == qn("w:tbl"):
            locator = f"T{ti + 1:03d}"
            table = doc.tables[ti]
            ti += 1
            source_block = source_blocks[locator]
            path = source_block["target"].split("#")[0]
            if path not in pages:
                title, kind = page_metadata(path, original["map"])
                pages[path] = {"path": path, "title": title, "kind": kind, "notice": NOTICE, "source_date": "2026-07-29", "blocks": []}
            page = pages[path]
            meta = tables[locator]
            rows = [[website_text(cell.text) for cell in row.cells] for row in table.rows]
            rows_html = [["<br>".join(paragraph_html(p._p, doc.part.rels) for p in cell.paragraphs) for cell in row.cells] for row in table.rows]
            block = {"type": "table", "source": locator, "id": "table-" + locator.lower(), "role": meta["role"], "caption": meta["caption"], "rows": rows, "rows_html": rows_html, "source_target": meta["target"]}
            block["presentation"] = "callout" if locator in {"T002", "T004", "T006", "T010", "T017"} else "contents" if locator == "T003" else "image-layout" if locator == "T022" else "table"
            block["is_empty_template"] = locator == "T016"
            page["blocks"].append(block)
            emitted_tables.append(locator)
        else:
            continue
        # Drawing-only body paragraphs are retained through image blocks. Drawing
        # objects inside T022 are processed here too, in their original order.
        for blip in element.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            if not rid or rid not in doc.part.rels:
                continue
            name = Path(str(doc.part.rels[rid].target_part.partname)).name
            figure = dict(figure_by_part[name])
            figure["source"] = locator
            page["blocks"].append(figure)
            emitted_images.append(figure["source_image"])

    expected = original["summary"]
    assert len(nonempty_paragraphs) == expected["paragraphs_nonempty"] == 437
    assert len(emitted_tables) == expected["tables"] == 23
    assert len(emitted_images) == expected["embedded_images"] == 9
    assert len(references) == 32
    assert len(nonempty_paragraphs) == len(set(nonempty_paragraphs))
    assert len(emitted_tables) == len(set(emitted_tables))
    assert len(emitted_images) == len(set(emitted_images))
    for page in pages.values():
        identifiers = [b["id"] for b in page["blocks"]]
        assert len(identifiers) == len(set(identifiers)), page["path"]
    summary = {
        "source_filename": source.name, "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "source_date": "2026-07-29", "paragraphs_total": pi,
        "paragraphs_nonempty_preserved": len(nonempty_paragraphs), "empty_paragraphs_omitted": len(empty_paragraphs),
        "tables_preserved": len(emitted_tables), "image_placements_preserved": len(emitted_images),
        "references_preserved": len(references), "page_count": len(pages),
        "coverage": {"nonempty_paragraph_ids": nonempty_paragraphs, "table_ids": emitted_tables, "image_ids": emitted_images, "empty_paragraph_ids": empty_paragraphs},
        "method": "نقل محتوى فقرات DOCX وجداوله وصوره في ترتيب المستند، مع حفظ التأكيد والروابط وبناء معرّفات HTML ثابتة، دون تحديث حالات الإنجاز الواردة في نسخة يوليو. استُبدلت الفاصلة المنقوطة العربية بالفاصلة العربية في نسخة الموقع وفق طلب الباحث.",
    }
    data = {"pages": list(pages.values()), "summary": summary}
    content = destination / "content"
    content.mkdir(exist_ok=True)
    (content / "source.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for page in data["pages"]:
        if page["kind"] == "chapter":
            md = content / "chapters" / (page["path"].split("/")[1] + ".md")
        elif page["kind"] == "appendix":
            md = content / "appendices" / (page["path"].split("/")[1] + ".md")
        elif page["path"] in {"about/index.html", "abstract/index.html", "references/index.html"}:
            md = content / "pages" / (page["path"].split("/")[0] + ".md")
        else:
            # Renderer consumes all pages from JSON. These frontmatter fragments
            # are retained in JSON, while the homepage/evidence layouts are built
            # separately from the navigable inventory.
            continue
        write_markdown(page, md)
    return data


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--inventory", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    data = extract(args.source.resolve(), args.inventory.resolve(), args.output.resolve())
    summary = {key: value for key, value in data["summary"].items() if key != "coverage"}
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
